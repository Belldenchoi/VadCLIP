"""Export RESULTS.md to Word with editable Office Math equations."""

from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "RESULTS.md"
OUTPUT = ROOT / "RESULTS.docx"
MML2OMML = Path(
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
)
MATH_NS = "http://www.w3.org/1998/Math/MathML"


def m(tag: str, *children, text: str | None = None):
    node = etree.Element(f"{{{MATH_NS}}}{tag}")
    if text is not None:
        node.text = text
    for child in children:
        # Formula symbols are intentionally reused across expressions. lxml
        # otherwise moves an existing node away from its previous parent.
        node.append(deepcopy(child))
    return node


def mi(text: str):
    return m("mi", text=text)


def mn(text: str):
    return m("mn", text=text)


def mo(text: str):
    return m("mo", text=text)


def mtext(text: str):
    return m("mtext", text=text)


def row(*items):
    return m("mrow", *items)


def sub(base, index):
    return m("msub", base, index)


def frac(numerator, denominator):
    return m("mfrac", numerator, denominator)


def fenced(content, open_char="(", close_char=")"):
    node = m("mfenced", content)
    node.set("open", open_char)
    node.set("close", close_char)
    return node


def norm(content):
    return row(mo("‖"), content, mo("‖"), sub(mi(""), mn("2")))


def function(name: str, argument):
    return row(mi(name), fenced(argument))


def sum_expr(index_text: str, condition, expression):
    return row(
        m("munderover", mo("∑"), row(mi(index_text), mo("="), condition), mi("N")),
        expression,
    )


def math_root(expression):
    return m("math", expression)


def formula_embedding():
    eci = sub(mi("e"), row(mi("c"), mo(","), mi("i")))
    dci = sub(mi("d"), row(mi("c"), mo(","), mi("i")))
    encoded = function("ftext", dci)
    return math_root(row(eci, mo("="), frac(encoded, norm(function("ftext", dci)))))


def formula_leave_one_out():
    qci = sub(mi("q"), row(mi("c"), mo(","), mi("i")))
    ecj = sub(mi("e"), row(mi("c"), mo(","), mi("j")))
    summation = row(
        m("munder", mo("∑"), row(mi("j"), mo("≠"), mi("i"))), ecj
    )
    mean = frac(summation, row(mi("N"), mo("−"), mn("1")))
    return math_root(row(qci, mo("="), function("norm", mean)))


def formula_centroid():
    muc = sub(mi("μ"), mi("c"))
    ecj = sub(mi("e"), row(mi("c"), mo(","), mi("j")))
    summation = m(
        "munderover", mo("∑"), row(mi("j"), mo("="), mn("1")), mi("N")
    )
    mean = frac(row(summation, ecj), mi("N"))
    return math_root(row(muc, mo("="), function("norm", mean)))


def formula_scores():
    eci = sub(mi("e"), row(mi("c"), mo(","), mi("i")))
    ac = sub(mi("a"), mi("c"))
    qci = sub(mi("q"), row(mi("c"), mo(","), mi("i")))
    muk = sub(mi("μ"), mi("k"))
    sname = sub(mi("s"), mtext("name"))
    sintra = sub(mi("s"), mtext("intra"))
    sinter = sub(mi("s"), mtext("inter"))
    stotal = sub(mi("s"), mtext("total"))
    maximum = row(
        m("munder", mi("max"), row(mi("k"), mo("≠"), mi("c"))),
        eci, mo("·"), muk,
    )
    return [
        math_root(row(sname, mo("="), eci, mo("·"), ac)),
        math_root(row(sintra, mo("="), eci, mo("·"), qci)),
        math_root(row(sinter, mo("="), maximum)),
        math_root(
            row(
                stotal, mo("="), mi("α"), sname, mo("+"), mi("β"), sintra,
                mo("−"), mi("γ"), sinter, mo(","), mtext("  "),
                mi("α"), mo("="), mn("0.4"), mo(","), mtext("  "),
                mi("β"), mo("="), mn("0.3"), mo(","), mtext("  "),
                mi("γ"), mo("="), mn("0.3"),
            )
        ),
    ]


def formula_diversity():
    eci = sub(mi("e"), row(mi("c"), mo(","), mi("i")))
    ecj = sub(mi("e"), row(mi("c"), mo(","), mi("j")))
    maximum = row(
        m("munder", mi("max"), row(mi("j"), mo("∈"), sub(mi("S"), mi("c")))),
        eci, mo("·"), ecj,
    )
    return math_root(row(mtext("accept "), eci, mtext(" if "), maximum, mo("<"), mi("τ")))


def formula_mean():
    pc = sub(mi("p"), mi("c"))
    sci = sub(mi("S"), mi("c"))
    eci = sub(mi("e"), row(mi("c"), mo(","), mi("i")))
    summation = row(m("munder", mo("∑"), row(mi("i"), mo("∈"), sci)), eci)
    average = frac(summation, row(mo("|"), sci, mo("|")))
    return math_root(row(pc, mo("="), function("norm", average)))


def formula_weighted():
    wi = sub(mi("w"), mi("i"))
    si = sub(mi("s"), mi("i"))
    sj = sub(mi("s"), mi("j"))
    numerator = m("msup", mi("e"), frac(si, mi("T")))
    denominator = row(
        m("munder", mo("∑"), row(mi("j"), mo("∈"), mi("S"))),
        m("msup", mi("e"), frac(sj, mi("T"))),
    )
    pc = sub(mi("p"), mi("c"))
    eci = sub(mi("e"), row(mi("c"), mo(","), mi("i")))
    weighted_sum = row(
        m("munder", mo("∑"), row(mi("i"), mo("∈"), sub(mi("S"), mi("c")))),
        wi, eci,
    )
    return [
        math_root(row(wi, mo("="), frac(numerator, denominator), mo(","), mtext("  "), mi("T"), mo("="), mn("0.10"))),
        math_root(row(pc, mo("="), function("norm", weighted_sum))),
    ]


def formula_logits():
    logit = sub(mi("logit2"), row(mi("t"), mo(","), mi("c")))
    vt = sub(mi("v"), mi("t"))
    pc = sub(mi("p"), mi("c"))
    cosine = row(function("norm", vt), mo("·"), pc)
    tau = sub(mi("τ"), mtext("logit"))
    return math_root(
        row(logit, mo("="), frac(cosine, tau), mo(","), mtext("  "), tau, mo("="), mn("0.07"))
    )


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def add_hyperlink(paragraph, text: str, target: str):
    relationship = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([props, text_element])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^]]+\]\([^)]+\))")


def add_inline(paragraph, value: str, bold: bool = False):
    cursor = 0
    for match in INLINE_PATTERN.finditer(value):
        if match.start() > cursor:
            run = paragraph.add_run(value[cursor:match.start()])
            run.bold = bold
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(120, 45, 45)
            run.bold = bold
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            label, target = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, target)
        cursor = match.end()
    if cursor < len(value):
        run = paragraph.add_run(value[cursor:])
        run.bold = bold


def add_equation(document: Document, transform, expression):
    mathml = etree.tostring(expression)
    omml = transform(etree.fromstring(mathml)).getroot()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph._p.append(parse_xml(etree.tostring(omml)))


def add_formula_group(document: Document, transform, key: str):
    if key == "scores":
        document.add_heading("Các công thức dùng để chấm điểm", level=4)
        document.add_paragraph("Chuẩn hóa embedding của description:")
        add_equation(document, transform, formula_embedding())
        document.add_paragraph("Tạo tham chiếu nội lớp và centroid của từng lớp:")
        add_equation(document, transform, formula_leave_one_out())
        add_equation(document, transform, formula_centroid())
        document.add_paragraph("Ba thành phần score và score xếp hạng cuối cùng:")
        for expression in formula_scores():
            add_equation(document, transform, expression)
    elif key == "diversity":
        document.add_paragraph("Điều kiện nhận một candidate khi bật diversity:")
        add_equation(document, transform, formula_diversity())
    elif key == "aggregation":
        document.add_heading("Các công thức gộp prototype", level=4)
        document.add_paragraph("Mean aggregation:")
        add_equation(document, transform, formula_mean())
        document.add_paragraph("Weighted-mean aggregation:")
        for expression in formula_weighted():
            add_equation(document, transform, expression)
    elif key == "logits":
        document.add_heading("Công thức logits của A-branch", level=4)
        add_equation(document, transform, formula_logits())


def configure_document(document: Document):
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (
        ("Title", 20, "17365D"),
        ("Heading 1", 15, "17365D"),
        ("Heading 2", 13, "2F5597"),
        ("Heading 3", 12, "4472C4"),
        ("Heading 4", 11.5, "5B9BD5"),
    ):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    if "Code Block" not in document.styles:
        code_style = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(9)

    add_page_number(section.footer.paragraphs[0])
    document.core_properties.title = "Kết quả thí nghiệm Direct Class Prototype trên UCF-Crime"
    document.core_properties.subject = "Class Prototype ablation results"


def add_table(document: Document, rows: list[list[str]]):
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = 1
    for row_index, values in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = values[column_index] if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, value, bold=row_index == 0)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    document.add_paragraph()


def parse_table(lines: list[str], start: int):
    raw_rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw_rows.append(lines[index].strip())
        index += 1
    rows = []
    for row_index, raw in enumerate(raw_rows):
        cells = [cell.strip() for cell in raw.strip("|").split("|")]
        if row_index == 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows, index


def export():
    if not MML2OMML.exists():
        raise FileNotFoundError(f"Office Math converter not found: {MML2OMML}")
    transform = etree.XSLT(etree.parse(str(MML2OMML)))
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)

    paragraph_buffer: list[str] = []

    def flush_paragraph():
        if paragraph_buffer:
            paragraph = document.add_paragraph()
            add_inline(paragraph, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer.clear()

    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                paragraph = document.add_paragraph(style="Code Block")
                paragraph.add_run("\n".join(code_lines))
                code_lines.clear()
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if stripped.startswith("### Bước 4"):
            flush_paragraph()
            add_formula_group(document, transform, "scores")
        if stripped.startswith("### Bước 5"):
            flush_paragraph()
            add_formula_group(document, transform, "diversity")
        if stripped.startswith("### Bước 6"):
            flush_paragraph()
            add_formula_group(document, transform, "aggregation")
        if stripped.startswith("### Danh sách class"):
            flush_paragraph()
            add_formula_group(document, transform, "logits")

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = heading.group(2)
            if level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, title)
            else:
                document.add_heading(title, level=level - 1)
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines):
            flush_paragraph()
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if bullet:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet.group(1))
            index += 1
            continue

        numbered = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if numbered:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(1))
            index += 1
            continue

        if not stripped:
            flush_paragraph()
        else:
            paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()
    document.save(OUTPUT)
    print(f"saved: {OUTPUT}")


if __name__ == "__main__":
    export()
